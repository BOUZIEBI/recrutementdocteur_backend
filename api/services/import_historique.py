from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile
from django.db import transaction
from rest_framework.exceptions import ValidationError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

import xlrd

from api.models import Historique


COLONNES = [
    "numero_candidat",
    "nom_prenoms",
    "sexe",
    "date_naissance",
    "lieu_naissance",
    "type_recrutement",
    "diplome",
    "domaine",
    "discipline",
    "specialite",
    "moyenne",
    "classement",
    "avis_cnres",
    "atructure_affectation",
    "observation",
    "situation_matrimoniale",
    "institution_dorigine",
    "date_obtention",
    "ecole_universite",
    "pays_obtention",
    "nombre_publication",
    "telephone",
    "handicap",
    "matricule"
]

CORRESPONDANCES_COLONNES = {
    "numero candidat": "numero_candidat",
    "nom prenoms": "nom_prenoms",
    "sexe": "sexe",
    "date de naissance": "date_naissance",
    "lieu de naissance": "lieu_naissance",
    "matricule": "matricule",
    #"date obtention": "date_obtention",
    "date obtention diplome": "date_obtention",
    "nombre de publication": "nombre_publication",
    "telephone": "telephone",
    "handicap": "handicap",
    "structure affectation": "structure_affectation",
    "type de recrutement": "type_recrutement",
    "type de concours": "type_concours",
    "domaine reesao": "domaine_reesao",
    "matricule": "matricule",
    "discipline": "discipline",
    "specialite": "specialite",
    "situation matrimoniale": "situation_matrimoniale",
    "ministere et institution": "ministere_institution",
    "diplome":"diplome",
    "date d'obtention": "date_obtention",
    "ecole ou universite ayant delivre le diplome": "ecole_universite",
    "pays obtention": "pays_obtention",
    "nombre de publiction": "nombre_publication‚‚",
}


from datetime import date, datetime, time

from django.db import transaction
from django.utils import timezone
from rest_framework.exceptions import ValidationError
import re
import unicodedata
from api.models import Historique



def normaliser_entete(valeur):
    if valeur is None:
        return ""

    texte = str(valeur).strip().lower()

    texte = unicodedata.normalize("NFKD", texte)
    texte = "".join(
        caractere
        for caractere in texte
        if not unicodedata.combining(caractere)
    )

    texte = texte.replace("_", " ")
    texte = re.sub(r"[^\w\s]", " ", texte)
    texte = re.sub(r"\s+", " ", texte).strip()

    return texte


def convertir_nom_colonne(entete: object) -> str | None:
    """
    Convertit un intitulé Excel en nom canonique.

    Retourne None si la colonne n'est pas reconnue.
    """

    nom = normaliser_entete(entete)

    if not nom:
        return None

    # Colonnes dont le titre peut contenir une session
    if (
        nom.startswith("numero candidat session")
        or nom.startswith("numero de candidat session")
    ):
        return "numero_candidat"

    if nom.startswith("nom prenoms") or nom.startswith("nom et prenoms"):
        return "nom_prenoms"

    # Correspondances personnalisées
    if nom in CORRESPONDANCES_COLONNES:
        return CORRESPONDANCES_COLONNES[nom]

    # Reconnaître les noms déjà présents dans COLONNES
    for colonne in COLONNES:
        if normaliser_entete(colonne) == nom:
            return colonne

    # Colonne inconnue
    return None



def nettoyer_valeur(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value or None

    return value


def nettoyer_texte(value):
    value = nettoyer_valeur(value)

    if value is None:
        return None

    return str(value)


def convertir_entier(value):
    value = nettoyer_valeur(value)

    if value is None:
        return None

    try:
        nombre = int(value)

        if float(value) != nombre or nombre < 0:
            raise ValueError()

        return nombre

    except (TypeError, ValueError, OverflowError):
        raise ValueError(
            f"Nombre de publications invalide : {value}"
        )


def convertir_date(value):
    value = nettoyer_valeur(value)

    if value is None:
        return None

    if isinstance(value, datetime):
        resultat = value

    elif isinstance(value, date):
        resultat = datetime.combine(value, time.min)

    elif isinstance(value, str):
        resultat = None

        for format_date in (
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d-%m-%Y",
            "%Y-%m-%d %H:%M:%S",
        ):
            try:
                resultat = datetime.strptime(value, format_date)
                break
            except ValueError:
                continue

        if resultat is None:
            raise ValueError(f"Date invalide : {value}")

    else:
        raise ValueError(
            f"Date Excel non reconnue : {value}"
        )

    if timezone.is_naive(resultat):
        resultat = timezone.make_aware(
            resultat,
            timezone.get_current_timezone()
        )

    return resultat


def convertir_booleen(value):
    value = nettoyer_valeur(value)

    if value is None:
        return None

    if isinstance(value, bool):
        return value

    texte = str(value).strip().lower()

    if texte in ("oui", "o", "true", "1", "vrai"):
        return True

    if texte in ("non", "n", "false", "0", "faux"):
        return False

    raise ValueError(
        f"Valeur de handicap invalide : {value}"
    )



def importer_historiques(
    fichier,
    annee_concours,
    type_session
):

    lignes = lire_fichier_excel(fichier)

    if not lignes:
        raise ValidationError({
            "fichier": "Le fichier Excel est vide."
        })

    entetes, index_entetes = trouver_entetes_excel(lignes)

    # Récupérer les numéros déjà présents dans PostgreSQL
    numeros_existants = set(
        Historique.objects
        .exclude(numero_candidat__isnull=True)
        .values_list("numero_candidat", flat=True)
    )

    # Contient aussi les numéros rencontrés pendant cet import
    numeros_connus = numeros_existants.copy()

    historiques_a_creer = []
    erreurs = []

    nombre_ignores = 0
    nombre_doublons_base = 0
    nombre_doublons_fichier = 0

    for numero_ligne, valeurs in enumerate(
        lignes[index_entetes + 1:],
        start=index_entetes + 2
    ):

        donnees_excel = dict(zip(entetes, valeurs))

        donnees = {
            colonne: nettoyer_valeur(
                donnees_excel.get(colonne)
            )
            for colonne in COLONNES
        }

        # Ignorer les lignes entièrement vides
        if all(valeur is None for valeur in donnees.values()):
            continue

        try:
            numero_candidat = nettoyer_texte(
                donnees.get("numero_candidat")
            )

            if not numero_candidat:
                raise ValueError(
                    "Le numéro du candidat est obligatoire."
                )

            # Vérifier les doublons dans PostgreSQL
            if numero_candidat in numeros_existants:
                nombre_doublons_base += 1
                nombre_ignores += 1
                continue

            # Vérifier les doublons dans le fichier Excel
            if numero_candidat in numeros_connus:
                nombre_doublons_fichier += 1
                nombre_ignores += 1
                continue

            historique = Historique(
                numero_candidat=numero_candidat,

                matricule=nettoyer_texte(
                    donnees.get("matricule")
                ),
                
                nom_prenoms=nettoyer_texte(
                    donnees.get("nom_prenoms")
                ),

                date_naissance=convertir_date(
                    donnees.get("date_naissance")
                ),

                date_obtention_diplome=convertir_date(
                    donnees.get("date_obtention")
                ),

                telephone=nettoyer_texte(
                    donnees.get("telephone")
                ),

                nombre_publication=convertir_entier(
                    donnees.get("nombre_publication")
                ),

                annee_concours=annee_concours,
                type_session=type_session,
            )

            historique.full_clean()

            historiques_a_creer.append(historique)

            # Empêcher un doublon dans le même fichier
            numeros_connus.add(numero_candidat)

        except (ValueError, ValidationError) as exc:
            erreurs.append({
                "ligne": numero_ligne,
                "message": str(exc)
            })

    if erreurs:
        raise ValidationError({
            "message": "Certaines lignes sont invalides.",
            "total_erreurs": len(erreurs),
            "erreurs": erreurs,
        })

    with transaction.atomic():
        Historique.objects.bulk_create(
            historiques_a_creer,
            batch_size=500
        )

    return {
        "success": True,
        "message": "Importation terminée avec succès.",
        "nombre_enregistres": len(historiques_a_creer),
        "nombre_ignores": nombre_ignores,
        "nombre_doublons_base": nombre_doublons_base,
        "nombre_doublons_fichier": nombre_doublons_fichier,
    }
    


def lire_fichier_excel(fichier) -> list[tuple]:
    """
    Lit un fichier Excel .xls ou .xlsx.

    Retourne :
        list[tuple] : toutes les lignes du fichier Excel.

    Les cellules vides sont converties en None.
    """

    extension = Path(fichier.name).suffix.lower()

    if extension not in (".xls", ".xlsx"):
        raise ValidationError({
            "fichier": (
                "Format non pris en charge. "
                "Utilisez un fichier .xls ou .xlsx."
            )
        })

    fichier.seek(0)

    lignes = []

    if extension == ".xlsx":

        workbook = None

        try:
            workbook = load_workbook(
                filename=BytesIO(fichier.read()),
                read_only=True,
                data_only=True
            )

            feuille = workbook.active

            if feuille is None:
                raise ValueError("Aucune feuille Excel trouvée.")

            for ligne in feuille.iter_rows(values_only=True):

                valeurs = tuple(
                    nettoyer_valeur(valeur)
                    for valeur in ligne
                )

                lignes.append(valeurs)

        except (
            InvalidFileException,
            BadZipFile,
            ValueError,
            OSError,
            EOFError,
            KeyError,
            TypeError,
        ) as exc:
            raise ValidationError({
                "fichier": "Le fichier .xlsx est invalide ou illisible."
            }) from exc

        finally:
            if workbook is not None:
                workbook.close()

    elif extension == ".xls":

        try:
            workbook = xlrd.open_workbook(
                file_contents=fichier.read()
            )

            feuille = workbook.sheet_by_index(0)

            for index in range(feuille.nrows):

                valeurs = []

                for cellule in feuille.row(index):

                    valeur = cellule.value

                    if cellule.ctype in (
                        xlrd.XL_CELL_EMPTY,
                        xlrd.XL_CELL_BLANK
                    ):
                        valeur = None

                    elif cellule.ctype == xlrd.XL_CELL_DATE:
                        valeur = xlrd.xldate_as_datetime(
                            valeur,
                            workbook.datemode
                        )

                    elif cellule.ctype == xlrd.XL_CELL_BOOLEAN:
                        valeur = bool(valeur)

                    elif cellule.ctype == xlrd.XL_CELL_NUMBER:
                        if float(valeur).is_integer():
                            valeur = int(valeur)

                    elif cellule.ctype == xlrd.XL_CELL_ERROR:
                        raise ValueError(
                            "Cellule Excel contenant une erreur "
                            f"à la ligne {index + 1}."
                        )

                    valeurs.append(
                        nettoyer_valeur(valeur)
                    )

                lignes.append(tuple(valeurs))

        except (
            xlrd.XLRDError,
            ValueError,
            OSError,
            IndexError,
            TypeError,
        ) as exc:
            raise ValidationError({
                "fichier": "Le fichier .xls est invalide ou illisible."
            }) from exc

    return lignes 



def trouver_entetes_excel(lignes):

    for index, ligne in enumerate(lignes[:30]):

        entetes = [
            convertir_nom_colonne(cellule)
            for cellule in ligne
        ]

        if (
            "numero_candidat" in entetes
            and "nom_prenoms" in entetes
        ):
            print(
                "Ligne des en-têtes :",
                index + 1,
                flush=True
            )

            print(
                "En-têtes normalisés :",
                entetes,
                flush=True
            )

            return entetes, index

    raise ValidationError({
        "fichier": "Impossible de trouver les en-têtes Excel."
    })